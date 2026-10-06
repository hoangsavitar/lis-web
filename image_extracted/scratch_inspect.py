import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'c:\Users\hoang\Documents\Code\lis\image_extracted\[LIS] Sản phẩm.xlsx', data_only=False)
print('Sheet names:', wb.sheetnames)

# Inspect Sheet 2
sheet = wb.worksheets[1]
print(f'\n--- Target Sheet: {sheet.title} ({sheet.max_row} rows, {sheet.max_column} cols) ---')

for r in range(1, min(100, sheet.max_row + 1)):
    row_vals = []
    has_val = False
    for c in range(1, min(20, sheet.max_column + 1)):
        cell = sheet.cell(r, c)
        val = cell.value
        hl = cell.hyperlink.target if cell.hyperlink else None
        if hl:
            val = f"{val} <LINK: {hl}>"
        if val is not None and str(val).strip() != "":
            has_val = True
        row_vals.append(val)
    if has_val:
        # print non-None or formatted
        formatted = [f"C{c}:{repr(v)}" for c, v in enumerate(row_vals, 1) if v is not None and str(v).strip() != ""]
        print(f"Row {r:2d}: {' | '.join(formatted)}")
